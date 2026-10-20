# Gera "Roteiro do Operador.docx": o roteiro original com cada deixa (ENTER)
# escrita nas linhas de anotação e entre as falas.
# Uso: python3 app/ferramentas/roteiro_operador.py   (na pasta da peça)
import copy, json, os, re, sys
import docx
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

RAIZ = sys.argv[1] if len(sys.argv) > 1 else os.getcwd()
ORIGINAL = os.path.join(RAIZ, 'O_Magico_de_Oz_VERSAO_PROFESSOR.docx')
SAIDA = os.path.join(RAIZ, 'Roteiro do Operador.docx')
dados = json.load(open(os.path.join(RAIZ, 'deixas.json'), encoding='utf-8'))

NOMES = {
    'fazenda': 'fazenda (vento leve e bichos)', 'vaca': 'vaca mugindo', 'galo': 'galo cantando', 'galinhas': 'galinhas',
    'vento_forte': 'vento forte', 'tornado': 'rugido do tornado', 'trovao_1': 'trovão', 'trovao_2': 'trovão', 'trovao_3': 'trovão',
    'trovao_longe_1': 'trovão distante', 'trovao_longe_2': 'trovão distante', 'chuva': 'chuva', 'porta': 'porta batendo',
    'casa_cai': 'casa caindo (estrondo)', 'magia_1': 'brilho mágico', 'magia_2': 'brilho mágico', 'magia_3': 'brilho mágico',
    'magia_vupt': 'redemoinho mágico', 'corvos': 'corvos', 'corvos_voam': 'corvos voando', 'palha': 'palha mexendo', 'tombo': 'tombo',
    'metal_range_1': 'lata rangendo', 'metal_range_2': 'lata rangendo', 'oleo': 'almotolia de óleo', 'floresta_noite': 'floresta à noite',
    'leao_1': 'rugido do leão', 'leao_2': 'rugido do leão', 'fanfarra': 'fanfarra', 'estrondo': 'ronco grave', 'fogo': 'labareda',
    'risada_bruxa': 'gargalhada da bruxa', 'agua': 'balde de água', 'chiado': 'chiado (derretendo)', 'vapor': 'vapor',
    'maquina': 'máquina', 'maquina_para': 'máquina parando', 'cortina': 'cortina', 'salto': 'batida do sapato', 'coracao': 'coração batendo',
    'passaros': 'passarinhos', 'balao': 'balão', 'aplausos': 'aplausos'
}
EFEITOS = {
    'brilhos': 'brilhos', 'nevoa': 'névoa', 'tremeluzir': 'luz tremendo', 'fumaca': 'fumaça', 'coracao': 'coração brilhando',
    'raiosDeLuz': 'raios de luz', 'vagalumes': 'vaga-lumes', 'redemoinho': 'redemoinho', 'detritos': 'vento com folhas e palha',
    'poeira': 'poeira de luz', 'chamas': 'fogo nas tochas', 'bolha': 'bolha da fada', 'vapor': 'vapor', 'estrelas': 'estrelas',
    'anel': 'anel de luz', 'chuva': 'chuva', 'clarao': 'clarão'
}
TRANS = {'dissolver': 'aparece devagar', 'preto': 'passando pelo escuro', 'branco': 'passando por um clarão', 'corte': 'troca na hora',
         'redemoinho': 'girando', 'cortina': 'como cortina', 'revelar': 'como cortina sendo puxada', 'zoom': 'aproximando', 'magia': 'com brilho mágico'}

def nome_arq(a):
    b = os.path.splitext(os.path.basename(a))[0]
    return NOMES.get(b, b.replace('_', ' '))

def descrever(acoes):
    img, som, fx = [], [], []
    add = lambda l, t: t not in l and l.append(t)
    for a in acoes:
        if 'imagem' in a: add(img, f"{os.path.splitext(a['imagem'])[0]} ({TRANS.get(a.get('transicao', 'dissolver'), '')})")
        elif 'video' in a: add(img, 'vídeo da casa voando')
        elif 'filtro' in a: add(img, 'muda a luz da imagem')
        elif 'blackout' in a: add(img, 'tela preta' if a['blackout'] else 'a imagem volta')
        elif 'titulo' in a: add(img, f"título “{a['titulo']}”")
        elif 'sobrepor' in a: add(img, 'imagem extra: ' + os.path.splitext(os.path.basename(a['sobrepor']))[0].replace('_', ' '))
        elif 'musica' in a: add(som, f"♪ música {os.path.splitext(a['musica'])[0]}" + (f" (a partir de {int(a['inicio'])} s)" if a.get('inicio') else ''))
        elif 'musicaVolume' in a: add(som, '♪ música abaixa' if a['musicaVolume'] < 50 else '♪ música sobe')
        elif 'pararMusica' in a: add(som, '♪ música sai')
        elif 'ambiente' in a: add(som, nome_arq(a['ambiente']) + ' (contínuo)')
        elif 'pararAmbiente' in a: add(som, 'sons de fundo saem' if a['pararAmbiente'] == 'todos' else nome_arq(a['pararAmbiente']) + ' sai')
        elif 'som' in a: add(som, nome_arq(a['som']))
        elif 'efeito' in a: add(fx, EFEITOS.get(a['efeito'], a['efeito']))
        elif 'raios' in a: add(fx, 'raios e trovões')
        elif 'flash' in a: add(fx, 'clarão')
        elif 'tremor' in a: add(fx, 'tela treme')
        elif 'disparar' in a: add(fx, 'labaredas')
    return img, som, fx

# deixas numeradas, com a âncora no roteiro
lista = []
for it in dados['roteiro']:
    if 'deixa' in it:
        lista.append(it['deixa'])
for i, d in enumerate(lista):
    d['n'] = i + 1

# ---------- localizar as âncoras no docx ----------
doc = docx.Document(ORIGINAL)
body = doc.element.body
blocos = []  # ('p', Paragraph) | ('t', Table)
comecou = False
for el in body.iterchildren():
    tag = el.tag.split('}')[1]
    if tag == 'p':
        p = Paragraph(el, doc)
        if p.text.strip() == 'ROTEIRO': comecou = True
        blocos.append(('p', p, comecou))
    elif tag == 'tbl':
        blocos.append(('t', Table(el, doc), comecou))

# Liga cada grupo de deixas seguidas ao parágrafo do docx que vem logo depois dele.
# Se houver uma linha de anotação (tabela) logo antes desse parágrafo, a deixa vai nela;
# senão entra um parágrafo amarelo antes da fala.
ordem = [(k, b) for k, b, c in blocos if c]
pos = 0
def achar(txt):
    global pos
    for j in range(pos, len(ordem)):
        k, b = ordem[j]
        if k == 'p' and txt and txt in b.text:
            pos = j + 1
            return j
    return None

alvos = []      # (grupo, j do parágrafo seguinte | None, j do último parágrafo achado)
grupo = []
ultimo_j = None
for it in dados['roteiro']:
    if 'deixa' in it:
        grupo.append(it['deixa'])
        continue
    txt = it.get('texto') or it.get('rubrica') or it.get('cena') or it.get('nota')
    j = achar(txt[:40]) if txt else None
    if grupo and (j is not None or 'cena' not in it):
        alvos.append((grupo, j, ultimo_j))
        grupo = []
    if j is not None:
        ultimo_j = j
if grupo:
    alvos.append((grupo, None, ultimo_j))

def elem(b): return b._element if isinstance(b, Paragraph) else b._tbl
usadas = set()
final = []
for grupo, j, uj in alvos:
    if j is not None:
        k = j - 1
        while k >= 0 and ordem[k][0] == 'p' and not ordem[k][1].text.strip():
            k -= 1
        if k >= 0 and ordem[k][0] == 't' and k not in usadas:
            usadas.add(k)
            final.append((grupo, 'tabela', ordem[k][1]))
        else:
            final.append((grupo, 'antes', ordem[j][1]))
    else:
        final.append((grupo, 'depois', ordem[uj][1] if uj is not None else None))

# ---------- escrever ----------
AZUL = RGBColor(0x0B, 0x4F, 0x8A)
VERMELHO = RGBColor(0xB0, 0x1E, 0x1E)
CINZA = RGBColor(0x55, 0x55, 0x55)

def sombrear(p, cor='FFF3B8'):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), cor)
    pPr.append(shd)
    bdr = OxmlElement('w:pBdr')
    left = OxmlElement('w:left')
    left.set(qn('w:val'), 'single'); left.set(qn('w:sz'), '24'); left.set(qn('w:space'), '4'); left.set(qn('w:color'), 'E0A800')
    bdr.append(left)
    pPr.append(bdr)

def escrever_deixa(p, d, compacto=False):
    r = p.add_run(f"▶ ENTER {d['n']}  "); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = VERMELHO
    r = p.add_run(d['nome']); r.bold = True; r.font.size = Pt(11.5)
    img, som, fx = descrever(d['acoes'])
    linhas = []
    if d.get('quando'): linhas.append(('Quando: ', d['quando']))
    if img: linhas.append(('Imagem: ', '; '.join(img)))
    if som: linhas.append(('Som: ', '; '.join(som)))
    if fx: linhas.append(('Efeitos: ', '; '.join(fx)))
    for rot, txt in linhas:
        p.add_run().add_break()
        r = p.add_run(rot); r.bold = True; r.font.size = Pt(9.5); r.font.color.rgb = CINZA
        r = p.add_run(txt); r.font.size = Pt(9.5)

def novo_par_depois(ref_el):
    novo = OxmlElement('w:p')
    ref_el.addnext(novo)
    return Paragraph(novo, doc)

def novo_par_antes(ref_el):
    novo = OxmlElement('w:p')
    ref_el.addprevious(novo)
    return Paragraph(novo, doc)

for grupo, tipo, bloco in final:
    if tipo == 'tabela':
        tab = bloco
        # usa as linhas da tabela; cria mais se precisar
        while len(tab.rows) < len(grupo):
            novo = copy.deepcopy(tab.rows[-1]._tr)
            tab._tbl.append(novo)
        for i, row in enumerate(tab.rows):
            cel = row.cells[0]
            p = cel.paragraphs[0]
            for r in list(p.runs): r._r.getparent().remove(r._r)
            if i < len(grupo):
                escrever_deixa(p, grupo[i])
    elif tipo == 'antes':
        for d in grupo:
            p = novo_par_antes(elem(bloco))
            sombrear(p)
            p.paragraph_format.space_after = Pt(6)
            escrever_deixa(p, d)
    else:
        # depois da última fala achada (fim da peça); encadeia para manter a ordem
        if 'ref_fim' not in globals() or ref_fim is None:
            ref_fim = elem(bloco) if bloco is not None else elem([b for k, b, c in blocos if k == 'p' and b.text.strip()][-1])
        for d in grupo:
            p = novo_par_depois(ref_fim)
            sombrear(p)
            p.paragraph_format.space_before = Pt(6)
            escrever_deixa(p, d)
            ref_fim = p._p

# capa
for k, b, c in blocos:
    if k == 'p' and 'VERSÃO DO PROFESSOR' in b.text:
        for r in list(b.runs): r._r.getparent().remove(r._r)
        r = b.add_run('ROTEIRO DO OPERADOR — Mesa de Cena'); r.bold = True
        b.add_run().add_break()
        r = b.add_run('Cada “▶ ENTER” é o momento de apertar ENTER no programa. Backspace desfaz a última deixa; Esc abaixa todo o som.')
        r.font.size = Pt(10)
        break

doc.core_properties.title = 'O Mágico de Oz – Roteiro do Operador'
doc.save(SAIDA)
print('ok:', SAIDA, len(lista), 'deixas;', sum(len(g) for g, t, b in final if t == 'tabela'), 'nas linhas de anotação')
